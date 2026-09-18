# -*- coding: utf-8 -*-
"""
Markdown -> Word 转换器（课程设计报告专用）

用法：
    cd doc/课程设计报告
    python build_docx.py

行为：
    按 CHAPTERS 列表顺序合并各章 .md，生成
    doc/企业应用开发实践课程设计报告_生成版.docx

支持：
    标题 H1~H4、段落、粗体、行内代码、无序/有序列表、
    引用块、围栏代码块、Markdown 表格、水平线、图片。

设计说明：
    纯 python-docx 实现（环境无 pandoc）。中文字体通过 w:eastAsia 单独设置，
    否则 Word 里中文会退回默认字体、与西文不一致。
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

# ---------------------------------------------------------------- 配置

BASE = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(os.path.dirname(BASE))     # 项目根目录
OUTPUT = os.path.join(PROJECT_ROOT, 'doc',
                      '企业应用开发实践课程设计报告_生成版.docx')

CHAPTERS = [
    '00_封面与摘要.md',
    '01_绪论.md',
    '02_相关技术基础.md',
    '03_系统需求分析.md',
    '04_系统总体设计.md',
    '05_系统详细设计与实现.md',
    '06_系统测试与优化.md',
    '07_总结与展望.md',
    '08_参考文献与附录.md',
]

CN_FONT = '宋体'
CN_HEAD_FONT = '黑体'
EN_FONT = 'Times New Roman'
CODE_FONT = 'Consolas'

# ---------------------------------------------------------------- 字体工具


def set_run_font(run, cn=CN_FONT, en=EN_FONT, size=None,
                 bold=None, color=None):
    """设置 run 的中英文字体（中文必须单独设 w:eastAsia）"""
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), en)
    rfonts.set(qn('w:hAnsi'), en)
    rfonts.set(qn('w:eastAsia'), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def style_document(doc):
    """设置全局样式：正文宋体小四、标题黑体"""
    normal = doc.styles['Normal']
    normal.font.name = EN_FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)

    sizes = {'Heading 1': 18, 'Heading 2': 16,
             'Heading 3': 14, 'Heading 4': 13}
    for name, size in sizes.items():
        st = doc.styles[name]
        st.font.name = EN_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.append(rfonts)
        rfonts.set(qn('w:eastAsia'), CN_HEAD_FONT)


def shade(cell, hex_color):
    """给表格单元格加底纹"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def add_code_block(doc, lines):
    """代码块：等宽字体 + 灰色底纹"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), 'F2F2F2')
    p_pr.append(shd)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, cn=CODE_FONT, en=CODE_FONT, size=9)

# ---------------------------------------------------------------- 行内解析

INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')


def add_inline(paragraph, text, base_size=None, base_bold=None):
    """解析 **粗体** 与 `行内代码`"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=base_size, bold=True)
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, cn=CODE_FONT, en=CODE_FONT,
                         size=(base_size or 10.5) - 0.5)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=base_size, bold=base_bold)

# ---------------------------------------------------------------- 表格

def parse_table(lines):
    """把 markdown 表格行解析成 rows（已去掉分隔行）"""
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-{2,}:?', c or '-') for c in cells):
            continue                      # 分隔行
        rows.append(cells)
    return rows


def add_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncol)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(ncol):
            text = row[ci] if ci < len(row) else ''
            cell = cells[ci]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, text, base_size=10.5, base_bold=(ri == 0))
            if ri == 0:
                shade(cell, 'D9E2F3')
    doc.add_paragraph()

# ---------------------------------------------------------------- 主解析

def convert_markdown(doc, md_text, image_base):
    lines = md_text.split('\n')
    i = 0
    para_buf = []

    def flush_para():
        if para_buf:
            text = ' '.join(para_buf).strip()
            if text:
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Pt(24)   # 首行缩进2字符
                add_inline(p, text)
            para_buf.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # ---- 围栏代码块 ----
        if stripped.startswith('```'):
            flush_para()
            lang = stripped[3:].strip().lower()
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                block.append(lines[i])
                i += 1
            i += 1

            if lang == 'mermaid':
                # Word 无法渲染 Mermaid。这里保留源码 + 给出渲染指引，
                # 便于作者到 mermaid.live 渲染后替换为图片。
                tip = doc.add_paragraph()
                tip.paragraph_format.left_indent = Inches(0.25)
                add_inline(tip,
                           '【下图待渲染】将下方 Mermaid 源码粘贴至 '
                           'https://mermaid.live 生成图片后插入此处',
                           base_size=10)
                add_code_block(doc, block)
            else:
                add_code_block(doc, block)
            continue

        # ---- 表格 ----
        if stripped.startswith('|') and stripped.count('|') >= 2:
            flush_para()
            tbl = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                tbl.append(lines[i])
                i += 1
            add_table(doc, parse_table(tbl))
            continue

        # ---- 标题 ----
        m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if m:
            flush_para()
            level, title = len(m.group(1)), m.group(2).strip()
            h = doc.add_heading(level=level)
            run = h.add_run(title)
            set_run_font(run, cn=CN_HEAD_FONT,
                         size={1: 18, 2: 16, 3: 14, 4: 13}[level],
                         bold=True)
            i += 1
            continue

        # ---- 水平线 ----
        if re.fullmatch(r'-{3,}|\*{3,}|_{3,}', stripped):
            flush_para()
            p = doc.add_paragraph()
            p_pr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:color'), '999999')
            pbdr.append(bottom)
            p_pr.append(pbdr)
            i += 1
            continue

        # ---- 图片 ----
        m = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if m:
            flush_para()
            alt, path = m.group(1), m.group(2)
            full = os.path.normpath(os.path.join(image_base, path))
            if os.path.isfile(full):
                try:
                    doc.add_picture(full, width=Inches(5.8))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    p = doc.add_paragraph()
                    add_inline(p, f'[插图失败：{path}]')
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # 图片缺失时留占位，提示此处需要补图或截图
                add_inline(p, f'【此处待插入图片：{alt}（{path}）】',
                           base_size=10.5)
            i += 1
            continue

        # ---- 引用块 ----
        if stripped.startswith('>'):
            flush_para()
            quote = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                quote.append(lines[i].strip().lstrip('>').strip())
                i += 1
            text = ' '.join(x for x in quote if x)
            if text:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p_pr = p._p.get_or_add_pPr()
                pbdr = OxmlElement('w:pBdr')
                left = OxmlElement('w:left')
                left.set(qn('w:val'), 'single')
                left.set(qn('w:sz'), '18')
                left.set(qn('w:color'), '8FAADC')
                pbdr.append(left)
                p_pr.append(pbdr)
                add_inline(p, text, base_size=10.5)
            continue

        # ---- 无序列表 ----
        if re.match(r'^[-*+]\s+', stripped):
            flush_para()
            text = re.sub(r'^[-*+]\s+', '', stripped)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text)
            i += 1
            continue

        # ---- 有序列表 ----
        if re.match(r'^\d+[.)]\s+', stripped):
            flush_para()
            text = re.sub(r'^\d+[.)]\s+', '', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text)
            i += 1
            continue

        # ---- 空行 = 段落结束 ----
        if not stripped:
            flush_para()
            i += 1
            continue

        # ---- 普通文本（累积成段） ----
        para_buf.append(stripped)
        i += 1

    flush_para()


# ---------------------------------------------------------------- 入口

def main():
    if not os.path.isdir(BASE):
        print('目录不存在:', BASE)
        return 1

    doc = Document()
    style_document(doc)

    missing = []
    for name in CHAPTERS:
        path = os.path.join(BASE, name)
        if not os.path.isfile(path):
            missing.append(name)
            continue
        with open(path, 'r', encoding='utf-8') as f:
            md = f.read()
        # 每章另起一页
        doc.add_page_break()
        # 图片路径相对于【章节目录】解析：章节目录是 doc/课程设计报告/，
        # 因此 md 里写 ../images/xxx.png 即 doc/images/xxx.png
        convert_markdown(doc, md, image_base=BASE)
        print('  已合并 %-28s (%d 字符)' % (name, len(md)))

    if missing:
        print('\n[警告] 以下章节文件不存在，已跳过：')
        for m in missing:
            print('   -', m)

    doc.save(OUTPUT)
    size_kb = os.path.getsize(OUTPUT) / 1024
    print('\n生成成功: %s (%.1f KB)' % (OUTPUT, size_kb))
    return 0


if __name__ == '__main__':
    sys.exit(main())
